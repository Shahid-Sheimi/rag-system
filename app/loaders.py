"""File loaders. Each returns a list of (text, metadata) sections."""
from pathlib import Path
import pandas as pd
from pypdf import PdfReader
from docx import Document as DocxDocument
from .config import settings


def load_pdf(path: Path):
    reader = PdfReader(str(path))
    out = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            out.append((text, {"page": i}))
    # NOTE: scanned PDFs have no text layer -> add OCR (e.g. ocrmypdf / pytesseract) if needed.
    return out


def load_docx(path: Path):
    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return [("\n".join(parts), {})] if parts else []


def load_csv(path: Path):
    df = pd.read_csv(path, dtype=str, keep_default_na=False, sep=None,
                     engine="python", encoding_errors="replace")
    n = settings.csv_rows_per_chunk
    out = []
    for start in range(0, len(df), n):
        block = df.iloc[start:start + n]
        lines = []
        for _, row in block.iterrows():
            lines.append(" | ".join(f"{c}: {v}" for c, v in row.items() if str(v).strip() != ""))
        out.append(("\n".join(lines), {"rows": f"{start + 2}-{start + len(block) + 1}"}))  # +2 = header + 1-based
    return out


def load_text(path: Path):
    text = path.read_text(encoding="utf-8", errors="replace")
    return [(text, {})] if text.strip() else []


LOADERS = {".pdf": load_pdf, ".docx": load_docx, ".csv": load_csv, ".txt": load_text, ".md": load_text}


def load_file(path: Path):
    return LOADERS[path.suffix.lower()](path)
